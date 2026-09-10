from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# =========================================================
# Decimal Helpers
# =========================================================

ZERO = Decimal("0.00")


def to_decimal(
    value,
    default=ZERO,
):
    try:
        if (
            value is None
            or value == ""
        ):
            return default

        return Decimal(
            str(value)
        ).quantize(
            Decimal("0.01")
        )

    except (
        InvalidOperation,
        TypeError,
        ValueError,
    ):
        return default


# =========================================================
# GST Rules
# =========================================================
#
# Final / discounted unit price:
#
# <= 2499
# GST  = 5%
# CGST = 2.5%
# SGST = 2.5%
#
# >= 2500
# GST  = 18%
# CGST = 9%
# SGST = 9%
#
# GST is calculated on the actual selling / discounted
# unit price stored in OrderItem.unit_price.
# =========================================================

GST_THRESHOLD = Decimal("2500.00")

LOW_GST_RATE = Decimal("5.00")
LOW_CGST_RATE = Decimal("2.50")
LOW_SGST_RATE = Decimal("2.50")

HIGH_GST_RATE = Decimal("18.00")
HIGH_CGST_RATE = Decimal("9.00")
HIGH_SGST_RATE = Decimal("9.00")


def get_gst_rates(
    unit_price,
):
    price = to_decimal(
        unit_price
    )

    if price >= GST_THRESHOLD:
        return {
            "gst_rate":
                HIGH_GST_RATE,

            "cgst_rate":
                HIGH_CGST_RATE,

            "sgst_rate":
                HIGH_SGST_RATE,
        }

    return {
        "gst_rate":
            LOW_GST_RATE,

        "cgst_rate":
            LOW_CGST_RATE,

        "sgst_rate":
            LOW_SGST_RATE,
    }


def calculate_item_gst(
    unit_price,
    quantity,
):
    price = to_decimal(
        unit_price
    )

    try:
        qty = int(
            quantity or 0
        )

    except (
        TypeError,
        ValueError,
    ):
        qty = 0

    if qty < 0:
        qty = 0

    taxable_value = (
        price
        * Decimal(
            str(qty)
        )
    ).quantize(
        Decimal("0.01")
    )

    rates = get_gst_rates(
        price
    )

    cgst_amount = (
        taxable_value
        * rates[
            "cgst_rate"
        ]
        / Decimal("100")
    ).quantize(
        Decimal("0.01")
    )

    sgst_amount = (
        taxable_value
        * rates[
            "sgst_rate"
        ]
        / Decimal("100")
    ).quantize(
        Decimal("0.01")
    )

    gst_amount = (
        cgst_amount
        + sgst_amount
    ).quantize(
        Decimal("0.01")
    )

    value_with_gst = (
        taxable_value
        + gst_amount
    ).quantize(
        Decimal("0.01")
    )

    return {
        "unit_price":
            price,

        "quantity":
            qty,

        "taxable_value":
            taxable_value,

        "gst_rate":
            rates[
                "gst_rate"
            ],

        "cgst_rate":
            rates[
                "cgst_rate"
            ],

        "cgst_amount":
            cgst_amount,

        "sgst_rate":
            rates[
                "sgst_rate"
            ],

        "sgst_amount":
            sgst_amount,

        "gst_amount":
            gst_amount,

        "value_with_gst":
            value_with_gst,
    }


# =========================================================
# Generic Object Helpers
# =========================================================

def safe_text(
    value,
):
    if value is None:
        return ""

    return str(
        value
    ).strip()


def get_order_number(
    order,
):
    return safe_text(
        getattr(
            order,
            "order_number",
            None,
        )
        or getattr(
            order,
            "order_id",
            None,
        )
        or getattr(
            order,
            "id",
            "",
        )
    )


def get_customer_email(
    order,
):
    user = getattr(
        order,
        "user",
        None,
    )

    email = getattr(
        order,
        "email",
        None,
    )

    if email:
        return safe_text(
            email
        )

    if user:
        return safe_text(
            getattr(
                user,
                "email",
                "",
            )
        )

    return ""


def get_customer_name(
    order,
):
    name = getattr(
        order,
        "full_name",
        None,
    )

    if name:
        return safe_text(
            name
        )

    user = getattr(
        order,
        "user",
        None,
    )

    if user:
        try:
            full_name = (
                user
                .get_full_name()
                .strip()
            )

            if full_name:
                return full_name

        except (
            AttributeError,
            TypeError,
        ):
            pass

        return safe_text(
            getattr(
                user,
                "username",
                "",
            )
        )

    return ""


def get_full_address(
    order,
):
    try:
        address = (
            order.full_address
        )

        if callable(
            address
        ):
            address = (
                address()
            )

        if address:
            return safe_text(
                address
            )

    except (
        AttributeError,
        TypeError,
    ):
        pass

    parts = [
        getattr(
            order,
            "address_line_1",
            "",
        ),
        getattr(
            order,
            "address_line_2",
            "",
        ),
        getattr(
            order,
            "landmark",
            "",
        ),
        getattr(
            order,
            "city",
            "",
        ),
        getattr(
            order,
            "state",
            "",
        ),
        getattr(
            order,
            "postal_code",
            "",
        ),
        getattr(
            order,
            "country",
            "",
        ),
    ]

    return ", ".join(
        safe_text(
            part
        )
        for part
        in parts
        if safe_text(
            part
        )
    )


def get_order_items(
    order,
):
    items_relation = getattr(
        order,
        "items",
        None,
    )

    if items_relation is None:
        return []

    try:
        return list(
            items_relation.all()
        )

    except AttributeError:
        try:
            return list(
                items_relation
            )
        except TypeError:
            return []


# =========================================================
# Date Helper
# =========================================================

def format_datetime(
    value,
):
    if not value:
        return ""

    try:
        if timezone.is_aware(
            value
        ):
            value = (
                timezone.localtime(
                    value
                )
            )

        return value.strftime(
            "%d-%m-%Y %I:%M %p"
        )

    except (
        AttributeError,
        TypeError,
        ValueError,
    ):
        return safe_text(
            value
        )


# =========================================================
# Excel Column Definitions
# =========================================================

HEADERS = [
    "Order Number",
    "Order Date",

    "Customer Name",
    "Customer Email",
    "Customer Phone",
    "Alternate Phone",

    "Address",
    "City",
    "State",
    "Pincode",
    "Country",

    "Product Name",
    "Product SKU",
    "Variant SKU",
    "Color",
    "Size",

    "Quantity",

    "Final Unit Price",
    "Item Taxable Value",

    "GST Rate %",
    "CGST Rate %",
    "CGST Amount",
    "SGST Rate %",
    "SGST Amount",
    "Total GST",
    "Value Including GST",

    "Order Subtotal",
    "Order Discount",
    "Shipping Charge",
    "Order Tax",
    "Order Total",

    "Coupon Code",

    "Payment Method",
    "Payment Status",
    "Order Status",
]


# =========================================================
# Build Sales Report Workbook
# =========================================================

def build_sales_report_workbook(
    orders,
):
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = (
        "Sales Report"
    )

    # =====================================================
    # Title
    # =====================================================

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(
            HEADERS
        ),
    )

    title_cell = sheet.cell(
        row=1,
        column=1,
        value=(
            "Yuvon Design Hub - Sales Report"
        ),
    )

    title_cell.font = Font(
        bold=True,
        size=16,
    )

    title_cell.alignment = (
        Alignment(
            horizontal="center",
            vertical="center",
        )
    )

    sheet.row_dimensions[
        1
    ].height = 28

    # =====================================================
    # GST Rule Information
    # =====================================================

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=len(
            HEADERS
        ),
    )

    gst_rule_cell = (
        sheet.cell(
            row=2,
            column=1,
            value=(
                "GST Rule: Final unit price up to "
                "₹2,499 = 5% GST "
                "(CGST 2.5% + SGST 2.5%); "
                "₹2,500 and above = 18% GST "
                "(CGST 9% + SGST 9%). "
                "GST calculated on final discounted "
                "selling price."
            ),
        )
    )

    gst_rule_cell.font = Font(
        italic=True,
    )

    gst_rule_cell.alignment = (
        Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )
    )

    sheet.row_dimensions[
        2
    ].height = 35

    # =====================================================
    # Header Row
    # =====================================================

    header_row = 4

    for (
        column_index,
        header,
    ) in enumerate(
        HEADERS,
        start=1,
    ):
        cell = sheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

        cell.font = Font(
            bold=True,
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

        cell.alignment = (
            Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        )

    # =====================================================
    # Data Rows
    # =====================================================

    current_row = (
        header_row + 1
    )

    report_total_taxable = ZERO
    report_total_cgst = ZERO
    report_total_sgst = ZERO
    report_total_gst = ZERO
    report_total_value_with_gst = ZERO

    for order in orders:

        items = get_order_items(
            order
        )

        if not items:
            continue

        for item in items:

            calculation = (
                calculate_item_gst(
                    getattr(
                        item,
                        "unit_price",
                        ZERO,
                    ),
                    getattr(
                        item,
                        "quantity",
                        0,
                    ),
                )
            )

            report_total_taxable += (
                calculation[
                    "taxable_value"
                ]
            )

            report_total_cgst += (
                calculation[
                    "cgst_amount"
                ]
            )

            report_total_sgst += (
                calculation[
                    "sgst_amount"
                ]
            )

            report_total_gst += (
                calculation[
                    "gst_amount"
                ]
            )

            report_total_value_with_gst += (
                calculation[
                    "value_with_gst"
                ]
            )

            values = [
                get_order_number(
                    order
                ),

                format_datetime(
                    getattr(
                        order,
                        "created_at",
                        None,
                    )
                ),

                get_customer_name(
                    order
                ),

                get_customer_email(
                    order
                ),

                safe_text(
                    getattr(
                        order,
                        "phone",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "alternate_phone",
                        "",
                    )
                ),

                get_full_address(
                    order
                ),

                safe_text(
                    getattr(
                        order,
                        "city",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "state",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "postal_code",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "country",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        item,
                        "product_name",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        item,
                        "product_sku",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        item,
                        "variant_sku",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        item,
                        "color",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        item,
                        "size",
                        "",
                    )
                ),

                calculation[
                    "quantity"
                ],

                float(
                    calculation[
                        "unit_price"
                    ]
                ),

                float(
                    calculation[
                        "taxable_value"
                    ]
                ),

                float(
                    calculation[
                        "gst_rate"
                    ]
                ),

                float(
                    calculation[
                        "cgst_rate"
                    ]
                ),

                float(
                    calculation[
                        "cgst_amount"
                    ]
                ),

                float(
                    calculation[
                        "sgst_rate"
                    ]
                ),

                float(
                    calculation[
                        "sgst_amount"
                    ]
                ),

                float(
                    calculation[
                        "gst_amount"
                    ]
                ),

                float(
                    calculation[
                        "value_with_gst"
                    ]
                ),

                float(
                    to_decimal(
                        getattr(
                            order,
                            "subtotal",
                            ZERO,
                        )
                    )
                ),

                float(
                    to_decimal(
                        getattr(
                            order,
                            "discount_amount",
                            ZERO,
                        )
                    )
                ),

                float(
                    to_decimal(
                        getattr(
                            order,
                            "shipping_charge",
                            ZERO,
                        )
                    )
                ),

                float(
                    to_decimal(
                        getattr(
                            order,
                            "tax_amount",
                            ZERO,
                        )
                    )
                ),

                float(
                    to_decimal(
                        getattr(
                            order,
                            "total_amount",
                            ZERO,
                        )
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "coupon_code",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "payment_method",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "payment_status",
                        "",
                    )
                ),

                safe_text(
                    getattr(
                        order,
                        "status",
                        "",
                    )
                ),
            ]

            for (
                column_index,
                value,
            ) in enumerate(
                values,
                start=1,
            ):
                cell = sheet.cell(
                    row=current_row,
                    column=column_index,
                    value=value,
                )

                cell.alignment = (
                    Alignment(
                        vertical="top",
                        wrap_text=True,
                    )
                )

            current_row += 1

    # =====================================================
    # Summary
    # =====================================================

    current_row += 1

    summary_start = (
        current_row
    )

    summary_data = [
        (
            "Report Taxable Value",
            report_total_taxable,
        ),
        (
            "Total CGST",
            report_total_cgst,
        ),
        (
            "Total SGST",
            report_total_sgst,
        ),
        (
            "Total GST",
            report_total_gst,
        ),
        (
            "Value Including GST",
            report_total_value_with_gst,
        ),
    ]

    for (
        label,
        amount,
    ) in summary_data:

        sheet.cell(
            row=current_row,
            column=1,
            value=label,
        ).font = Font(
            bold=True,
        )

        amount_cell = (
            sheet.cell(
                row=current_row,
                column=2,
                value=float(
                    amount
                ),
            )
        )

        amount_cell.font = Font(
            bold=True,
        )

        amount_cell.number_format = (
            '₹#,##0.00'
        )

        current_row += 1

    # =====================================================
    # Currency Formatting
    # =====================================================

    currency_columns = [
        18,  # Final Unit Price
        19,  # Taxable Value
        22,  # CGST Amount
        24,  # SGST Amount
        25,  # Total GST
        26,  # Including GST
        27,  # Subtotal
        28,  # Discount
        29,  # Shipping
        30,  # Order Tax
        31,  # Order Total
    ]

    for column_number in (
        currency_columns
    ):
        for row_number in range(
            header_row + 1,
            summary_start,
        ):
            sheet.cell(
                row=row_number,
                column=column_number,
            ).number_format = (
                '₹#,##0.00'
            )

    # =====================================================
    # Percentage Formatting
    # =====================================================

    percentage_columns = [
        20,
        21,
        23,
    ]

    for column_number in (
        percentage_columns
    ):
        for row_number in range(
            header_row + 1,
            summary_start,
        ):
            sheet.cell(
                row=row_number,
                column=column_number,
            ).number_format = (
                '0.00'
            )

    # =====================================================
    # Freeze Header
    # =====================================================

    sheet.freeze_panes = (
        "A5"
    )

    # =====================================================
    # Auto Filter
    # =====================================================

    if current_row > (
        header_row + 1
    ):
        sheet.auto_filter.ref = (
            f"A{header_row}:"
            f"{get_column_letter(len(HEADERS))}"
            f"{max(header_row, summary_start - 2)}"
        )

    # =====================================================
    # Column Widths
    # =====================================================

    column_widths = {
        1: 20,
        2: 22,

        3: 24,
        4: 28,
        5: 16,
        6: 16,

        7: 45,
        8: 18,
        9: 18,
        10: 14,
        11: 15,

        12: 35,
        13: 18,
        14: 18,
        15: 14,
        16: 12,

        17: 10,

        18: 18,
        19: 20,

        20: 14,
        21: 14,
        22: 18,
        23: 14,
        24: 18,
        25: 18,
        26: 22,

        27: 18,
        28: 18,
        29: 18,
        30: 18,
        31: 18,

        32: 18,

        33: 18,
        34: 18,
        35: 18,
    }

    for (
        column_number,
        width,
    ) in (
        column_widths.items()
    ):
        sheet.column_dimensions[
            get_column_letter(
                column_number
            )
        ].width = width

    return workbook


# =========================================================
# Create XLSX HTTP Response
# =========================================================

def create_sales_report_response(
    orders,
    filename=None,
):
    workbook = (
        build_sales_report_workbook(
            orders
        )
    )

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(
        0
    )

    if not filename:
        filename = (
            "yuvon_sales_report_"
            + timezone.localdate().strftime(
                "%Y-%m-%d"
            )
            + ".xlsx"
        )

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response